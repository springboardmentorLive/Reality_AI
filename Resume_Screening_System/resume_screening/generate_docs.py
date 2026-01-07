from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

def create_document():
    doc = Document()
    
    # --- Styles ---
    # (Simple styling via code)
    
    # --- 1. Title Page ---
    # Vertical spacing
    for _ in range(5):
        doc.add_paragraph()
        
    title = doc.add_heading('PROJECT EXPLANATION DOCUMENT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Resume Screening & ATS System')
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    # Spacing
    for _ in range(10):
        doc.add_paragraph()
        
    details = doc.add_paragraph()
    details.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = details.add_run(f"Author: Mowleen\n")
    run.font.size = Pt(14)
    run = details.add_run(f"Domain: AI / Machine Learning / NLP\n")
    run.font.size = Pt(14)
    run = details.add_run(f"Date: {datetime.date.today().strftime('%B %d, %Y')}")
    run.font.size = Pt(14)
    
    doc.add_page_break()
    
    # --- 2. Abstract ---
    doc.add_heading('2. Abstract', level=1)
    p = doc.add_paragraph(
        "The Resume Screening & ATS System is a Machine Learning-powered automated tool designed "
        "to assist HR professionals and recruiters in efficiently validating and categorizing resumes. "
        "By leveraging Natural Language Processing (NLP) techniques, the system extracts text from PDF and TXT resumes, "
        "identifies key skills, and predicts the most suitable job role using a trained Logistic Regression model. "
        "Furthermore, it acts as an Applicant Tracking System (ATS) by comparing resumes against specific job descriptions "
        "to provide a relevance score and actionable feedback. This project aims to reduce the manual effort involved in "
        "initial resume screening and improve the objectivity of the recruitment process."
    )
    
    # --- 3. Introduction ---
    doc.add_heading('3. Introduction', level=1)
    doc.add_heading('3.1 Background & Motivation', level=2)
    doc.add_paragraph(
        "In the modern recruitment landscape, companies receive hundreds or thousands of resumes for a single job opening. "
        "Manually reviewing each document is time-consuming, prone to human error, and often biased. "
        "Recruiters need a way to quickly filter candidates based on relevant skills and experience."
    )
    doc.add_heading('3.2 Real-world Relevance', level=2)
    doc.add_paragraph(
        "Automated Resume Screening Systems are becoming standard in the industry. They allow organizations to "
        "focus their human resources on interviewing valid candidates rather than reading irrelevant CVs. "
        "This project simulates such a real-world enterprise tool."
    )
    
    # --- 4. Problem Statement ---
    doc.add_heading('4. Problem Statement', level=1)
    doc.add_paragraph(
        "The core problem addressed by this project is the inefficiency of manual resume screening. "
        "Traditional methods are unscalable and inconsistent. Specifically, the challenges include:\n"
        "- High volume of applications leading to fatigue.\n"
        "- Inconsistent evaluation criteria across different recruiters.\n"
        "- Difficulty in quickly matching long text documents (resumes) with specific keyword requirements."
    )
    
    # --- 5. Proposed Solution ---
    doc.add_heading('5. Proposed Solution', level=1)
    doc.add_paragraph(
        "We propose an AI-driven web application that automates the screening process. The solution works as follows:"
    )
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Input:").bold = True
    p.add_run(" User uploads a resume (PDF/TXT) and optionally provides a Job Description.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Processing:").bold = True
    p.add_run(" The system cleans the text, extracting only relevant alphanumeric content.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Analysis:").bold = True
    p.add_run(" TF-IDF vectorization converts text to numbers. A Logistic Regression model predicts the job category (e.g., Data Scientist, Java Developer).")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Output:").bold = True
    p.add_run(" A dashboard displaying the predicted role, confidence score, ATS match percentage, and missing skills.")
    
    # --- 6. System Architecture ---
    doc.add_heading('6. System Architecture', level=1)
    doc.add_paragraph("The system consists of the following key components:")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Frontend (Streamlit):").bold = True
    p.add_run(" Provides an interactive user interface for file upload and visualization.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Backend Logic (Python):").bold = True
    p.add_run(" Handles file parsing, text cleaning, and business logic.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Machine Learning Model:").bold = True
    p.add_run(" A trained classifier (Logistic Regression) powered by Scikit-Learn.")

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Data Understanding:").bold = True
    p.add_run(" Uses TF-IDF (Term Frequency-Inverse Document Frequency) to understand the importance of words.")

    # --- 7. Technologies Used ---
    doc.add_heading('7. Technologies Used', level=1)
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Technology'
    hdr_cells[1].text = 'Purpose'
    
    techs = [
        ("Python", "Core programming language for backend and logic."),
        ("Streamlit", "Web framework for building the UI quickly."),
        ("Scikit-Learn", "Machine Learning library for model training and TF-IDF."),
        ("Pandas", "Data manipulation and handling datasets."),
        ("NLTK / Regex", "Natural Language Processing (Text Cleaning)."),
        ("Plotly", "Interactive charts and visualizations (Gauge, Bar charts).")
    ]
    
    for tech, desc in techs:
        row_cells = table.add_row().cells
        row_cells[0].text = tech
        row_cells[1].text = desc

    # --- 8. Key Features ---
    doc.add_heading('8. Key Features', level=1)
    features = [
        ("Resume Parsing", "Extracts readable text from PDF and TXT files automatically."),
        ("Role Prediction", "Classifies resumes into categories like 'Data Scientist', 'HR', 'Web Developer'."),
        ("ATS Scoring", "Calculates a match percentage between the resume and a provided Job Description."),
        ("Skill Extraction", "Identifies specific technical skills present in the resume text."),
        ("Visual Analytics", "Displays confidence scores and keyword gaps using interactive charts."),
        ("Deep Dark UI", "A modern, professional dark-themed user interface.")
    ]
    for feat, expl in features:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(feat + ": ").bold = True
        p.add_run(expl)

    # --- 9. Advantages ---
    doc.add_heading('9. Advantages', level=1)
    doc.add_paragraph(style='List Bullet').add_run("Speed: Drastically reduces the time needed to screen candidates.")
    doc.add_paragraph(style='List Bullet').add_run("Objectivity: Removes unconscious bias by relying on keyword and text patterns.")
    doc.add_paragraph(style='List Bullet').add_run("Scalability: Can process any number of resumes instantly.")
    doc.add_paragraph(style='List Bullet').add_run("User-Friendly: Simple interface requires no technical knowledge to use.")

    # --- 10. Limitations ---
    doc.add_heading('10. Limitations / Disadvantages', level=1)
    doc.add_paragraph(style='List Bullet').add_run("Formatting Sensitivity: Complex resume layouts (tables, graphics) might lead to poor text extraction.")
    doc.add_paragraph(style='List Bullet').add_run("Keyword Dependency: The model heavily relies on specific keywords; synonyms might sometimes be missed if not trained.")
    doc.add_paragraph(style='List Bullet').add_run("Category Limit: Can only predict roles that were present in the training dataset.")

    # --- 11. Weaknesses ---
    doc.add_heading('11. Weaknesses', level=1)
    doc.add_paragraph(
        "The system interprets text literally. It may not understand context or 'soft skills' as well as a human."
        " Additionally, very short or poorly written resumes may yield low confidence scores even if the candidate is capable."
    )

    # --- 12. Future Enhancements ---
    doc.add_heading('12. Future Enhancements', level=1)
    doc.add_paragraph(style='List Bullet').add_run("Deep Learning Integration: Move from TF-IDF/Logistic Regression to BERT or Transformers for semantic understanding.")
    doc.add_paragraph(style='List Bullet').add_run("Experience Parsing: Automatically detect years of experience.")
    doc.add_paragraph(style='List Bullet').add_run("Multiple Page Analysis: Better handling of multi-page PDFs.")
    doc.add_paragraph(style='List Bullet').add_run("Database Integration: Save candidate profiles for future retrieval.")

    # --- 13. Applications ---
    doc.add_heading('13. Applications / Use Cases', level=1)
    doc.add_paragraph("- Recruitment Agencies handling bulk hiring.")
    doc.add_paragraph("- Corporate HR Departments for initial screening.")
    doc.add_paragraph("- Job Portals to recommend jobs to candidates.")
    doc.add_paragraph("- Educational Institutes to check student resume quality.")

    # --- 14. Conclusion ---
    doc.add_heading('14. Conclusion', level=1)
    doc.add_paragraph(
        "The Resume Screening System successfully demonstrates the application of Machine Learning in HR Tech. "
        "By automating the repetitive task of reading resumes, it allows recruiters to add value where it matters most—interviewing and "
        "assessing culture fit. This project not only solves a critical business problem but also showcases the power of NLP and Streamlit "
        "in building practical data applications."
    )

    # Save
    filename = 'Project_Explanation_Resume_Screening.docx'
    doc.save(filename)
    print(f"Document saved as {filename}")

if __name__ == "__main__":
    create_document()
