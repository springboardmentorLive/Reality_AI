
# import streamlit as st
# import re
# import pdfplumber
# from docx import Document
# import pandas as pd
# from sklearn.feature_extraction.text import TfidfVectorizer
# from sklearn.metrics.pairwise import cosine_similarity
# import plotly.graph_objects as go

# # ================= PAGE CONFIG =================
# st.set_page_config(
#     page_title="AI Resume Analyzer & ATS",
#     page_icon="📄",
#     layout="wide"
# )

# # ================= UI STYLES =================
# st.markdown("""
# <style>
# .stApp { background-color:#0b1220; color:#e5e7eb; }

# .card {
#     background:#020617;
#     padding:22px;
#     border-radius:14px;
#     box-shadow:0 4px 14px rgba(0,0,0,0.6);
#     margin-bottom:24px;
# }

# .skill-box {
#     display:inline-block;
#     background:#111827;
#     border:1px solid #374151;
#     padding:6px 14px;
#     border-radius:20px;
#     margin:6px 6px 0 0;
#     font-size:13px;
# }

# .skill-good {
#     background:#064e3b;
#     border:1px solid #22c55e;
# }

# .skill-miss {
#     background:#7f1d1d;
#     border:1px solid #f87171;
# }

# .good { background:#064e3b; padding:14px; border-radius:10px; }
# .mid { background:#78350f; padding:14px; border-radius:10px; }
# .bad { background:#7f1d1d; padding:14px; border-radius:10px; }

# h1 { font-size:28px; }
# h2 { font-size:22px; }
# h3 { font-size:18px; }
# </style>
# """, unsafe_allow_html=True)

# # ================= SIDEBAR =================
# st.sidebar.markdown("""
# <div style="text-align:center">
#     <img src="https://cdn-icons-png.flaticon.com/512/3135/3135715.png" width="60">
#     <h2>AI Resume ATS</h2>
#     <p style="font-size:13px;">ATS-based Resume Analysis</p>
# </div>
# <hr>
# """, unsafe_allow_html=True)

# module = st.sidebar.radio(
#     "Choose Module",
#     ["🧠 Employee – Resume Analyzer", "🏢 Employer – Resume Screening"]
# )

# # ================= HEADER =================
# st.markdown("""
# <div class="card">
# <img src="https://images.unsplash.com/photo-1504384308090-c894fdcc538d"
#  style="width:100%;max-height:200px;object-fit:cover;border-radius:12px;">
# <h1 style="text-align:center;margin-top:10px;">AI Resume Analyzer & ATS</h1>
# <p style="text-align:center;">ML & NLP Based Recruitment Assistant</p>
# </div>
# """, unsafe_allow_html=True)

# # ================= ROLE SKILLS =================
# ROLE_SKILLS = {
#     "AI / ML Engineer": [
#         "python","machine learning","deep learning","tensorflow",
#         "pytorch","nlp","statistics","data preprocessing"
#     ],
#     "Backend Developer": [
#         "python","java","node","django","flask",
#         "spring","api","mysql","postgresql"
#     ],
#     "Data Analyst": [
#         "python","sql","excel","power bi",
#         "tableau","data analysis","statistics"
#     ],
#     "Data Scientist": [
#         "python","machine learning","statistics","sql",
#         "pandas","numpy","scikit-learn"
#     ],
#     "Web Developer": [
#         "html","css","javascript","react","node"
#     ],
#     "Accountant": [
#         "accounting","tally","gst","taxation",
#         "financial statements","auditing"
#     ],
#     "HR Executive": [
#         "recruitment","hr operations","employee relations",
#         "payroll","onboarding"
#     ],
#     "Digital Marketing Executive": [
#         "seo","social media","content marketing",
#         "google ads","analytics"
#     ],
#     "Business Analyst": [
#         "business analysis","requirement gathering",
#         "stakeholder management","sql","documentation"
#     ]
# }

# # ================= HELPERS =================
# def read_pdf(file):
#     text = ""
#     with pdfplumber.open(file) as pdf:
#         for page in pdf.pages:
#             text += page.extract_text() or ""
#     return text

# def read_docx(file):
#     doc = Document(file)
#     return " ".join(p.text for p in doc.paragraphs)

# def clean_text(text):
#     text = text.lower()
#     text = re.sub(r"[^a-zA-Z0-9 ]", " ", text)
#     return re.sub(r"\s+", " ", text)

# def predict_roles(text):
#     roles = list(ROLE_SKILLS.keys())
#     corpus = [" ".join(ROLE_SKILLS[r]) for r in roles] + [text]
#     tfidf = TfidfVectorizer()
#     vectors = tfidf.fit_transform(corpus)
#     sims = cosine_similarity(vectors[-1], vectors[:-1])[0]
#     return dict(zip(roles, sims))

# def extract_skills(text, role):
#     found = [s for s in ROLE_SKILLS[role] if s in text]
#     missing = list(set(ROLE_SKILLS[role]) - set(found))
#     return found, missing

# def calculate_ats(found):
#     return round(min(len(found) * 7 + 40, 100), 1)

# # ================= VISUALS =================
# def ats_gauge(score):
#     fig = go.Figure(go.Indicator(
#         mode="gauge+number",
#         value=score,
#         title={"text":"ATS Score"},
#         gauge={
#             "axis":{"range":[0,100]},
#             "bar":{"color":"#22c55e"},
#             "steps":[
#                 {"range":[0,40],"color":"#7f1d1d"},
#                 {"range":[40,70],"color":"#ca8a04"},
#                 {"range":[70,100],"color":"#166534"}
#             ]
#         }
#     ))
#     fig.update_layout(height=360)
#     st.plotly_chart(fig, use_container_width=True)

# def keyword_donut(found, missing):
#     fig = go.Figure(go.Pie(
#         labels=["Matched","Missing"],
#         values=[len(found),len(missing)],
#         hole=0.6
#     ))
#     fig.update_layout(height=360)
#     st.plotly_chart(fig, use_container_width=True)

# def role_confidence_bar(scores):
#     fig = go.Figure(go.Bar(
#         x=list(scores.keys()),
#         y=[round(v*100,1) for v in scores.values()],
#         marker_color="#38bdf8"
#     ))
#     fig.update_layout(height=320, yaxis_title="Confidence %")
#     st.plotly_chart(fig, use_container_width=True)

# # ================= EMPLOYEE MODULE =================
# if module == "🧠 Employee – Resume Analyzer":

#     resume = st.file_uploader("📂 Upload Resume", ["pdf","docx","txt"])
#     analyze = st.button("🚀 Analyze Resume", use_container_width=True)

#     if resume and analyze:
#         raw = read_pdf(resume) if resume.name.endswith(".pdf") else read_docx(resume)
#         clean = clean_text(raw)

#         role_scores = predict_roles(clean)
#         target_role = max(role_scores, key=role_scores.get)
#         found, missing = extract_skills(clean, target_role)
#         ats = calculate_ats(found)

#         st.markdown('<div class="card">', unsafe_allow_html=True)
#         st.markdown("### 📄 Resume Preview")
#         st.text_area("", raw, height=320)
#         st.markdown('</div>', unsafe_allow_html=True)

#         st.markdown('<div class="card">', unsafe_allow_html=True)
#         ats_gauge(ats)
#         st.markdown('</div>', unsafe_allow_html=True)

#         st.markdown('<div class="card">', unsafe_allow_html=True)
#         st.markdown("### 🎯 Role Prediction Confidence")
#         role_confidence_bar(role_scores)
#         st.markdown('</div>', unsafe_allow_html=True)

#         st.markdown('<div class="card">', unsafe_allow_html=True)
#         st.markdown("### 🛠 Skill Analysis")

#         st.markdown("**Matched Skills**")
#         for s in found:
#             st.markdown(f"<span class='skill-box skill-good'>{s}</span>", unsafe_allow_html=True)

#         st.markdown("<br>**Missing Skills**", unsafe_allow_html=True)
#         for s in missing:
#             st.markdown(f"<span class='skill-box skill-miss'>{s}</span>", unsafe_allow_html=True)

#         st.markdown('</div>', unsafe_allow_html=True)

#         st.markdown('<div class="card">', unsafe_allow_html=True)
#         st.markdown("### 🍩 Keyword Match Ratio")
#         keyword_donut(found, missing)
#         st.markdown('</div>', unsafe_allow_html=True)

#         st.markdown('<div class="card">', unsafe_allow_html=True)
#         st.markdown("### ✅ Final Recommendation")

#         if ats >= 80 and len(missing) <= 2:
#             st.markdown("<div class='good'>Excellent resume. You are well-positioned for this role.</div>", unsafe_allow_html=True)
#         elif ats >= 60:
#             st.markdown("<div class='mid'>Good resume. Improve missing skills to increase selection chances.</div>", unsafe_allow_html=True)
#         else:
#             st.markdown("<div class='bad'>Resume needs improvement. Focus on skill alignment and ATS structure.</div>", unsafe_allow_html=True)

#         st.markdown('</div>', unsafe_allow_html=True)

# # ================= EMPLOYER MODULE =================
# if module == "🏢 Employer – Resume Screening":

#     jd = st.text_area("🧾 Paste Job Description")
#     resumes = st.file_uploader("📂 Upload Resumes", ["pdf","docx"], accept_multiple_files=True)
#     run = st.button("🔍 Run Screening", use_container_width=True)

#     if jd and resumes and run:
#         jd_clean = clean_text(jd)
#         resume_texts = [
#             clean_text(read_pdf(r) if r.name.endswith(".pdf") else read_docx(r))
#             for r in resumes
#         ]

#         tfidf = TfidfVectorizer()
#         vectors = tfidf.fit_transform([jd_clean] + resume_texts)

#         results = []
#         for i, r in enumerate(resumes):
#             score = cosine_similarity(vectors[0], vectors[i+1])[0][0] * 100
#             results.append({
#                 "Candidate": r.name,
#                 "Match Score (%)": round(score,2),
#                 "Decision": "Shortlisted" if score >= 55 else "Rejected"
#             })

#         st.markdown('<div class="card">', unsafe_allow_html=True)
#         st.markdown("### 📋 Screening Results")
#         st.dataframe(
#             pd.DataFrame(results).sort_values("Match Score (%)", ascending=False),
#             use_container_width=True
#         )
#         st.markdown('</div>', unsafe_allow_html=True)

import streamlit as st
import re
import pdfplumber
from docx import Document
import pandas as pd
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import plotly.graph_objects as go

# ================= PAGE CONFIG =================
st.set_page_config(
    page_title="AI Resume Analyzer & ATS",
    page_icon="📄",
    layout="wide"
)

# ================= UI STYLES =================
st.markdown("""
<style>
.stApp { background-color:#0b1220; color:#e5e7eb; }

.card {
    background:#020617;
    padding:22px;
    border-radius:14px;
    box-shadow:0 4px 14px rgba(0,0,0,0.6);
    margin-bottom:24px;
}

.skill-box {
    display:inline-block;
    background:#111827;
    border:1px solid #374151;
    padding:6px 14px;
    border-radius:20px;
    margin:6px 6px 0 0;
    font-size:13px;
}

.skill-good { background:#064e3b; border:1px solid #22c55e; }
.skill-miss { background:#7f1d1d; border:1px solid #f87171; }

.good { background:#064e3b; padding:14px; border-radius:10px; }
.mid { background:#78350f; padding:14px; border-radius:10px; }
.bad { background:#7f1d1d; padding:14px; border-radius:10px; }

h1 { font-size:28px; }
h2 { font-size:22px; }
h3 { font-size:18px; }
</style>
""", unsafe_allow_html=True)

# ================= SIDEBAR =================
st.sidebar.markdown("""
<div style="text-align:center">
    <img src="https://cdn-icons-png.flaticon.com/512/3135/3135715.png" width="60">
    <h2>AI Resume ATS</h2>
    <p style="font-size:13px;">ATS-based Resume Analysis</p>
</div>
<hr>
""", unsafe_allow_html=True)

module = st.sidebar.radio(
    "Choose Module",
    ["🧠 Employee – Resume Analyzer", "🏢 Employer – Resume Screening"]
)

# ================= HEADER =================
st.markdown("""
<div class="card">
<img src="https://images.unsplash.com/photo-1504384308090-c894fdcc538d"
 style="width:100%;max-height:200px;object-fit:cover;border-radius:12px;">
<h1 style="text-align:center;margin-top:10px;">AI Resume Analyzer & ATS</h1>
<p style="text-align:center;">ML & NLP Based Recruitment Assistant</p>
</div>
""", unsafe_allow_html=True)

# ================= ROLE SKILLS =================
ROLE_SKILLS = {
    "AI / ML Engineer": [
        "python","machine learning","deep learning","tensorflow",
        "pytorch","nlp","statistics","data preprocessing"
    ],
    "Backend Developer": [
        "python","java","node","django","flask",
        "spring","api","mysql","postgresql"
    ],
    "Data Analyst": [
        "python","sql","excel","power bi",
        "tableau","data analysis","statistics"
    ],
    "Data Scientist": [
        "python","machine learning","statistics","sql",
        "pandas","numpy","scikit-learn"
    ],
    "Web Developer": [
        "html","css","javascript","react","node"
    ],
    "Accountant": [
        "accounting","tally","gst","taxation",
        "financial statements","auditing"
    ],
    "HR Executive": [
        "recruitment","hr operations","employee relations",
        "payroll","onboarding"
    ],
    "Digital Marketing Executive": [
        "seo","social media","content marketing",
        "google ads","analytics"
    ],
    "Business Analyst": [
        "business analysis","requirement gathering",
        "stakeholder management","sql","documentation"
    ]
}

# ================= HELPERS =================
def read_pdf(file):
    text = ""
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            text += page.extract_text() or ""
    return text

def read_docx(file):
    doc = Document(file)
    return " ".join(p.text for p in doc.paragraphs)

def clean_text(text):
    text = text.lower()
    text = re.sub(r"[^a-zA-Z0-9 ]", " ", text)
    return re.sub(r"\s+", " ", text)

def predict_roles(text):
    roles = list(ROLE_SKILLS.keys())
    corpus = [" ".join(ROLE_SKILLS[r]) for r in roles] + [text]
    tfidf = TfidfVectorizer()
    vectors = tfidf.fit_transform(corpus)
    sims = cosine_similarity(vectors[-1], vectors[:-1])[0]
    return dict(zip(roles, sims))

def extract_skills(text, role):
    found = [s for s in ROLE_SKILLS[role] if s in text]
    missing = list(set(ROLE_SKILLS[role]) - set(found))
    return found, missing

def calculate_ats(found):
    return round(min(len(found) * 7 + 40, 100), 1)

def selection_probability(ats, role_conf, skill_ratio):
    prob = (ats * 0.5) + (role_conf * 100 * 0.3) + (skill_ratio * 100 * 0.2)
    return round(min(prob, 95), 1)

# ================= VISUALS =================
def ats_gauge(score):
    fig = go.Figure(go.Indicator(
        mode="gauge+number",
        value=score,
        title={"text":"ATS Score"},
        gauge={
            "axis":{"range":[0,100]},
            "bar":{"color":"#22c55e"},
            "steps":[
                {"range":[0,40],"color":"#7f1d1d"},
                {"range":[40,70],"color":"#ca8a04"},
                {"range":[70,100],"color":"#166534"}
            ]
        }
    ))
    fig.update_layout(height=360)
    st.plotly_chart(fig, use_container_width=True)

def keyword_donut(found, missing):
    fig = go.Figure(go.Pie(
        labels=["Matched","Missing"],
        values=[len(found),len(missing)],
        hole=0.6
    ))
    fig.update_layout(height=360)
    st.plotly_chart(fig, use_container_width=True)

def role_confidence_bar(scores):
    fig = go.Figure(go.Bar(
        x=list(scores.keys()),
        y=[round(v*100,1) for v in scores.values()],
        marker_color="#38bdf8"
    ))
    fig.update_layout(height=320, yaxis_title="Confidence %")
    st.plotly_chart(fig, use_container_width=True)

# ================= EMPLOYEE MODULE =================
if module == "🧠 Employee – Resume Analyzer":

    resume = st.file_uploader("📂 Upload Resume", ["pdf","docx","txt"])
    analyze = st.button("🚀 Analyze Resume", use_container_width=True)

    if resume and analyze:
        raw = read_pdf(resume) if resume.name.endswith(".pdf") else read_docx(resume)
        clean = clean_text(raw)

        role_scores = predict_roles(clean)
        matched_role = max(role_scores, key=role_scores.get)
        role_conf = role_scores[matched_role]

        found, missing = extract_skills(clean, matched_role)
        ats = calculate_ats(found)
        skill_ratio = len(found) / max(len(found) + len(missing), 1)

        selection_prob = selection_probability(ats, role_conf, skill_ratio)

        # Resume Preview
        st.markdown('<div class="card">', unsafe_allow_html=True)
        st.markdown("### 📄 Resume Preview")
        st.text_area("", raw, height=320)
        st.markdown('</div>', unsafe_allow_html=True)

        # Matched Role + Selection Probability
        st.markdown('<div class="card">', unsafe_allow_html=True)
        c1, c2 = st.columns(2)
        c1.metric("🎯 Matched Role", matched_role)
        c2.metric("📈 Selection Probability", f"{selection_prob}%")
        st.markdown('</div>', unsafe_allow_html=True)

        # ATS
        st.markdown('<div class="card">', unsafe_allow_html=True)
        ats_gauge(ats)
        st.markdown('</div>', unsafe_allow_html=True)

        # Role confidence
        st.markdown('<div class="card">', unsafe_allow_html=True)
        st.markdown("### 🎯 Role Prediction Confidence")
        role_confidence_bar(role_scores)
        st.markdown('</div>', unsafe_allow_html=True)

        # Skill analysis
        st.markdown('<div class="card">', unsafe_allow_html=True)
        st.markdown("### 🛠 Skill Analysis")

        st.markdown("**Matched Skills**")
        for s in found:
            st.markdown(f"<span class='skill-box skill-good'>{s}</span>", unsafe_allow_html=True)

        st.markdown("<br>**Missing Skills**", unsafe_allow_html=True)
        for s in missing:
            st.markdown(f"<span class='skill-box skill-miss'>{s}</span>", unsafe_allow_html=True)

        st.markdown('</div>', unsafe_allow_html=True)

        # Keyword donut
        st.markdown('<div class="card">', unsafe_allow_html=True)
        st.markdown("### 🍩 Keyword Match Ratio")
        keyword_donut(found, missing)
        st.markdown('</div>', unsafe_allow_html=True)

        # Final Recommendation
        st.markdown('<div class="card">', unsafe_allow_html=True)
        st.markdown("### ✅ Final Recommendation")

        if selection_prob >= 80:
            st.markdown("<div class='good'>Strong profile. High chance of shortlisting.</div>", unsafe_allow_html=True)
        elif selection_prob >= 60:
            st.markdown("<div class='mid'>Good profile. Improve missing skills.</div>", unsafe_allow_html=True)
        else:
            st.markdown("<div class='bad'>Needs improvement to pass ATS screening.</div>", unsafe_allow_html=True)

        st.markdown('</div>', unsafe_allow_html=True)

# ================= EMPLOYER MODULE =================
if module == "🏢 Employer – Resume Screening":

    jd = st.text_area("🧾 Paste Job Description")
    resumes = st.file_uploader("📂 Upload Resumes", ["pdf","docx"], accept_multiple_files=True)
    run = st.button("🔍 Run Screening", use_container_width=True)

    if jd and resumes and run:
        jd_clean = clean_text(jd)
        resume_texts = [
            clean_text(read_pdf(r) if r.name.endswith(".pdf") else read_docx(r))
            for r in resumes
        ]

        tfidf = TfidfVectorizer()
        vectors = tfidf.fit_transform([jd_clean] + resume_texts)

        results = []
        for i, r in enumerate(resumes):
            score = cosine_similarity(vectors[0], vectors[i+1])[0][0] * 100
            results.append({
                "Candidate": r.name,
                "Match Score (%)": round(score,2),
                "Decision": "Shortlisted" if score >= 30 else "Rejected"
            })

        st.markdown('<div class="card">', unsafe_allow_html=True)
        st.markdown("### 📋 Screening Results")
        st.dataframe(
            pd.DataFrame(results).sort_values("Match Score (%)", ascending=False),
            use_container_width=True
        )
        st.markdown('</div>', unsafe_allow_html=True)

